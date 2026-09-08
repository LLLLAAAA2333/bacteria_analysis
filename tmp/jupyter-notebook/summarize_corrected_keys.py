from pathlib import Path
import re

import numpy as np
import pandas as pd
from Bio import Phylo
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data"

df = pd.read_parquet(DATA / "106bac.parquet")
df["sample_id"] = df["stim_name"].str.extract(r"^(A\d{3})", expand=False)

recording_key = ["date", "worm_key"]
trial_key = recording_key + ["sample_id", "segment_index"]
trace_key = trial_key + ["neuron"]
row_key = trace_key + ["time_point"]

trials = df[trial_key + ["stimulus", "stim_name"]].drop_duplicates()
traces = df[trace_key].drop_duplicates()

print("CORE")
print("rows", len(df))
print("recordings(date×worm)", df.groupby(recording_key).ngroups)
print("trials", len(trials))
print("traces", len(traces))
print("duplicate correct row key", df.duplicated(row_key).sum())
print("rows/trace values", df.groupby(trace_key).size().value_counts().sort_index().to_dict())
print("unique timepoints/trace", df.groupby(trace_key)["time_point"].nunique().value_counts().sort_index().to_dict())

print("\nRECORDINGS PER DATE")
print(df.groupby("date")["worm_key"].nunique().to_string())

print("\nTRIALS PER SAMPLE")
per_sample = trials.groupby("sample_id").agg(
    n_trials=("segment_index", "size"),
    n_recordings=("worm_key", "size"),
    n_dates=("date", "nunique"),
)
sample_recordings = trials[recording_key + ["sample_id"]].drop_duplicates()
per_sample["n_recordings"] = sample_recordings.groupby("sample_id").size()
print(per_sample.describe().round(3).to_string())
print("minimum examples")
print(per_sample.nsmallest(10, "n_trials").to_string())
print("maximum examples")
print(per_sample.nlargest(10, "n_trials").to_string())

print("\nSEGMENTS PER SAMPLE×RECORDING")
print(trials.groupby(recording_key + ["sample_id"]).size().value_counts().sort_index().to_string())

print("\nDATE SUMMARY")
date_summary = trials.groupby("date").agg(
    n_trials=("segment_index", "size"),
    n_samples=("sample_id", "nunique"),
    n_worms=("worm_key", "nunique"),
)
print(date_summary.to_string())

print("\nNEURON PANEL STABILITY")
trial_neuron_n = traces.groupby(trial_key).size()
print("neurons/trial summary")
print(trial_neuron_n.describe().round(3).to_string())

recording_union = traces.groupby(recording_key)["neuron"].agg(set)
recording_trials = trials.groupby(recording_key).size()
expected_pairs = sum(len(recording_union.loc[key]) * int(recording_trials.loc[key]) for key in recording_union.index)
observed_pairs = len(traces)
print("expected trace slots within recording-specific neuron union", expected_pairs)
print("observed traces", observed_pairs)
print("structural missing within recording panels", (expected_pairs - observed_pairs) / expected_pairs)

neuron_panel_stats = []
for neuron in sorted(df["neuron"].unique()):
    eligible_recordings = [key for key, panel in recording_union.items() if neuron in panel]
    eligible_trials = sum(int(recording_trials.loc[key]) for key in eligible_recordings)
    observed = int((traces["neuron"] == neuron).sum())
    neuron_panel_stats.append((neuron, observed, eligible_trials, 1 - observed / eligible_trials))
neuron_panel_stats = pd.DataFrame(
    neuron_panel_stats, columns=["neuron", "observed_trials", "eligible_trials", "missing_rate_within_recording_panel"]
).sort_values("missing_rate_within_recording_panel", ascending=False)
print(neuron_panel_stats.to_string(index=False, formatters={"missing_rate_within_recording_panel": "{:.3%}".format}))

sets_per_recording = (
    traces.groupby(trial_key)["neuron"].agg(lambda x: tuple(sorted(x)))
    .groupby(level=recording_key).nunique()
)
print("unique neuron sets per recording distribution", sets_per_recording.value_counts().sort_index().to_dict())

print("\nBASELINE")
baseline = df[df["time_point"] < df["start_time"]]
baseline_stats = baseline.groupby(trace_key)["delta_F_over_F0"].agg(
    baseline_mean="mean", baseline_sd="std", baseline_range=lambda x: x.max() - x.min()
)
trial_baseline = baseline_stats.groupby(level=trial_key).agg(
    median_neuron_sd=("baseline_sd", "median"),
    p90_neuron_sd=("baseline_sd", lambda x: x.quantile(0.9)),
    max_neuron_sd=("baseline_sd", "max"),
)
print("max abs baseline trace mean", baseline_stats["baseline_mean"].abs().max())
print("trial baseline summaries")
print(trial_baseline.describe(percentiles=[.5, .9, .95, .99]).round(4).to_string())
print("worst trials by median neuron baseline SD")
print(trial_baseline.nlargest(10, "median_neuron_sd").round(4).to_string())

print("\nEXTREMES")
trace_stats = df.groupby(trace_key)["delta_F_over_F0"].agg(trace_min="min", trace_max="max")
trace_stats["max_abs"] = trace_stats.abs().max(axis=1)
trace_stats = trace_stats.join(baseline_stats)
print("max abs quantiles")
print(trace_stats["max_abs"].quantile([.5, .9, .95, .99, .999, 1]).round(4).to_string())
for threshold in [2, 3, 5, 8]:
    extreme = trace_stats[trace_stats["max_abs"] >= threshold].reset_index()
    print(f">={threshold}: n={len(extreme)}, neurons={extreme['neuron'].value_counts().to_dict()}")
print("top 15 with Axxx")
top = trace_stats.nlargest(15, "max_abs").reset_index().merge(
    df[["stimulus", "sample_id"]].drop_duplicates(), on="sample_id", how="left"
)
print(top.round(4).to_string(index=False))

print("\nSOURCE SETS")
neural_ids = set(df["sample_id"].unique())
matrix = pd.read_excel(DATA / "matrix.xlsx", index_col=0)
matrix.index = matrix.index.astype(str).str.strip()
matrix_ids = set(matrix.index)
species = pd.read_excel(DATA / "GM300_bacteria_species_summary.xlsx", sheet_name="Axxx_species_mapping")
species_ids = set(species["AID"].dropna().astype(str).str.strip())
tree_ids = {t.name for t in Phylo.read(DATA / "16S.aln.trim.fa.treefile", "newick").get_terminals()}
wb = load_workbook(DATA / "metabolism_raw_data.xlsx", read_only=True)
header = [c.value for c in next(wb["all"].iter_rows(min_row=1, max_row=1))]
raw_ids = {str(v).strip() for v in header if re.fullmatch(r"A\d{3}", str(v).strip())}
wb.close()
for name, ids in {"neural": neural_ids, "matrix": matrix_ids, "species": species_ids, "tree": tree_ids, "raw": raw_ids}.items():
    print(name, len(ids))
print("matrix-species", sorted(matrix_ids - species_ids), sorted(species_ids - matrix_ids))
print("tree-species", sorted(tree_ids - species_ids), sorted(species_ids - tree_ids))
print("raw-only vs matrix", sorted(raw_ids - matrix_ids))
print("matrix-only vs raw", sorted(matrix_ids - raw_ids))

