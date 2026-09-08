from pathlib import Path

import nbformat as nbf


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "output" / "jupyter-notebook" / "01_measurement_first_inspection.ipynb"

nb = nbf.read(OUT, as_version=4)
cells = []


def md(text: str) -> None:
    cells.append(nbf.v4.new_markdown_cell(text))


def code(text: str) -> None:
    cells.append(nbf.v4.new_code_cell(text))


md(
    """# 01 — 从 measurement 开始认识数据

这个 notebook 只回答测量层面的第一批问题：数据的行粒度、样本与 trial 覆盖、日期结构、neuron panel、缺失、时间轴、baseline 波动和极端 $\Delta F/F_0$ traces。

本轮不做 RSA、HMDS、聚类、feature screen 或机制解释，也不读取 `backup/`。所有结论分为两类：

- **Observation**：可直接从现有文件确认。
- **Assumption / unresolved**：仅凭当前表无法确认，需要实验 protocol 或更上游数据。

Notebook 规则：每个分析代码 cell 前有 **Question**，执行后紧跟 **Observation**。"""
)

md(
    """## 当前生物学问题

在讨论“不同细菌如何被神经系统表征”之前，先确认我们实际测量了什么、独立观测单位是什么、哪些不完整或异常，以及刺激前后时间窗如何定义。

完成标准：

1. 给出可复用但仍属探索阶段的 measurement key。
2. 明确 106 个细菌的 trial/date/recording 覆盖。
3. 把三类 missingness 分开：数值缺失、timepoint 缺失、neuron panel 覆盖。
4. 明确 onset/offset 在表中的标记，并显式保留 protocol 未决项。
5. 展示 baseline 波动和极端 traces 的原始曲线，不自动剔除。"""
)

md("""### Question

当前环境能否只从仓库的 primary data 文件复现这一轮检查？""")

code(
    r'''from pathlib import Path
import re

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
from Bio import Phylo
from openpyxl import load_workbook

pd.set_option("display.max_columns", 40)
pd.set_option("display.max_rows", 120)
pd.set_option("display.max_colwidth", 120)
sns.set_theme(style="whitegrid", context="notebook")
plt.rcParams["figure.dpi"] = 110
plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]

# 允许从仓库根目录或 notebook 所在目录启动 kernel。
start = Path.cwd().resolve()
ROOT = next(
    (candidate for candidate in [start, *start.parents]
     if (candidate / "data" / "106bac.parquet").exists()),
    None,
)
if ROOT is None:
    raise FileNotFoundError("找不到 data/106bac.parquet；请从仓库内启动 notebook。")

DATA = ROOT / "data"
FILES = {
    "neural": DATA / "106bac.parquet",
    "lcms_fold_change": DATA / "matrix.xlsx",
    "lcms_raw": DATA / "metabolism_raw_data.xlsx",
    "taxonomy": DATA / "GM300_bacteria_species_summary.xlsx",
    "tree": DATA / "16S.aln.trim.fa.treefile",
    "current_samples": DATA / "current_samples.xlsx",
}

file_inventory = pd.DataFrame(
    {
        "exists": {name: path.exists() for name, path in FILES.items()},
        "size_MiB": {name: path.stat().st_size / 2**20 for name, path in FILES.items()},
    }
).round(2)
file_inventory'''
)

md(
    """### Observation

六个 primary data 文件都存在。路径解析不依赖隐藏的工作目录状态；本 notebook 不读取或修改 `backup/`，也不写回任何原始文件。"""
)

md("""### Question

`106bac.parquet` 的行、列、类型和最基本的数据质量是什么？""")

code(
    r'''df = pd.read_parquet(FILES["neural"])

basic_summary = pd.Series(
    {
        "rows": len(df),
        "columns": df.shape[1],
        "memory_MiB": df.memory_usage(deep=True).sum() / 2**20,
        "null_cells": int(df.isna().sum().sum()),
        "nonfinite_delta_F_over_F0": int((~np.isfinite(df["delta_F_over_F0"].to_numpy())).sum()),
        "unique_neurons": df["neuron"].nunique(),
        "unique_stimulus_codes": df["stimulus"].nunique(),
        "unique_dates": df["date"].nunique(),
        "unique_worm_labels": df["worm_key"].nunique(),
        "unique_time_points": df["time_point"].nunique(),
    },
    name="value",
)
display(basic_summary.to_frame())
display(df.dtypes.rename("dtype").to_frame())
display(df.head(8))'''
)

md(
    """### Observation

神经表有 2,191,545 行和 11 个显式字段；`delta_F_over_F0` 没有 NaN 或 ±inf。共有 28 个 neuron label、106 个 stimulus code、9 个日期、8 个重复使用的 `w1…w8` 标签，以及 45 个 timepoint。用户口述的“40 s”不能直接等同于表中的 40 个点，因为实际索引是 0–44。"""
)

md("""### Question

内部 `stimulus` code、`stim_name` 和 Axxx 样本号是否一一对应？""")

code(
    r'''df["sample_id"] = df["stim_name"].str.extract(r"^(A\d{3})", expand=False)

identity_checks = pd.Series(
    {
        "rows_without_Axxx": int(df["sample_id"].isna().sum()),
        "unique_sample_ids": df["sample_id"].nunique(),
        "max_Axxx_per_stimulus": int(df.groupby("stimulus")["sample_id"].nunique().max()),
        "max_stimulus_per_Axxx": int(df.groupby("sample_id")["stimulus"].nunique().max()),
        "unique_stim_name_suffixes": df["stim_name"].str.replace(
            r"^A\d{3}\s*", "", regex=True
        ).nunique(),
    },
    name="value",
)
display(identity_checks.to_frame())
display(
    df[["stimulus", "sample_id", "stim_name", "stim_color"]]
    .drop_duplicates()
    .sort_values("sample_id")
    .head(15)
)'''
)

md(
    """### Observation

106 个内部 `stimulus` code 与 106 个 Axxx 样本号严格一一对应，所有 `stim_name` 都以 `stationary` 结尾。后续以 `sample_id` 作为细菌身份，保留 `stimulus` 仅用于追溯原表。"""
)

md("""### Question

`worm_key` 是否能单独标识一只独立 worm？忽略日期会不会制造重复？""")

code(
    r'''wrong_row_key = ["worm_key", "sample_id", "segment_index", "neuron", "time_point"]
correct_row_key = ["date", *wrong_row_key]

key_check = pd.Series(
    {
        "duplicate_rows_without_date": int(df.duplicated(wrong_row_key).sum()),
        "duplicate_rows_with_date": int(df.duplicated(correct_row_key).sum()),
        "worm_labels": df["worm_key"].nunique(),
        "date_x_worm_recordings": df.groupby(["date", "worm_key"]).ngroups,
        "dates_per_worm_label_min": int(df.groupby("worm_key")["date"].nunique().min()),
        "dates_per_worm_label_max": int(df.groupby("worm_key")["date"].nunique().max()),
    },
    name="value",
)
display(key_check.to_frame())
display(pd.crosstab(df["date"], df["worm_key"]).gt(0).astype(int))'''
)

md(
    """### Observation

`worm_key` 不是跨日期唯一的动物 ID。忽略 `date` 会产生 1,530 个伪重复行；加入 `date` 后重复数为 0。当前最小可用的 recording key 是 `(date, worm_key)`，共有 49 个 recording。统计独立性仍需 protocol 确认这些 recording 是否各来自不同动物。"""
)

md("""### Question

从表的唯一性出发，trial、trace 和 timepoint 的 operational key 应如何定义？""")

code(
    r'''recording_key = ["date", "worm_key"]
trial_key = [*recording_key, "sample_id", "segment_index"]
trace_key = [*trial_key, "neuron"]
row_key = [*trace_key, "time_point"]

trials = df[trial_key + ["stimulus", "stim_name"]].drop_duplicates()
traces = df[trace_key].drop_duplicates()
trace_time = df.groupby(trace_key)["time_point"].agg(["size", "nunique", "min", "max"])

trial_contract = pd.Series(
    {
        "recordings": df.groupby(recording_key).ngroups,
        "operational_trials": len(trials),
        "neuron_traces": len(traces),
        "duplicate_rows_under_full_key": int(df.duplicated(row_key).sum()),
        "min_rows_per_trace": int(trace_time["size"].min()),
        "max_rows_per_trace": int(trace_time["size"].max()),
        "min_unique_timepoints_per_trace": int(trace_time["nunique"].min()),
        "max_unique_timepoints_per_trace": int(trace_time["nunique"].max()),
    },
    name="value",
)
display(trial_contract.to_frame())

segments_per_sample_recording = trials.groupby(
    [*recording_key, "sample_id"]
).size().value_counts().sort_index().rename("n_sample_recordings")
display(segments_per_sample_recording.to_frame().rename_axis("trials_per_sample_recording"))'''
)

md(
    """### Observation

用 `(date, worm_key, sample_id, segment_index)` 作为 operational trial key，可得到 2,696 个 trials；再加 `neuron` 得到 48,701 条 traces；再加 `time_point` 后每行唯一。每条 trace 都恰好有 45 个不同 timepoints（0–44）。同一 sample × recording 通常有 4 或 6 个 `segment_index`，因此 `segment_index` 很像重复呈现索引，但它的实验语义仍需 protocol 确认。"""
)

md("""### Question

神经、LCMS fold change、原始 LCMS、分类信息和 16S tree 的样本集合如何对齐？""")

code(
    r'''matrix = pd.read_excel(FILES["lcms_fold_change"], index_col=0)
matrix.index = matrix.index.astype(str).str.strip()
matrix_ids = set(matrix.index)

species = pd.read_excel(
    FILES["taxonomy"], sheet_name="Axxx_species_mapping"
)
species_ids = set(species["AID"].dropna().astype(str).str.strip())

tree = Phylo.read(FILES["tree"], "newick")
tree_ids = {terminal.name for terminal in tree.get_terminals() if terminal.name}

raw_book = load_workbook(FILES["lcms_raw"], read_only=True, data_only=False)
raw_header = [cell.value for cell in next(raw_book["all"].iter_rows(min_row=1, max_row=1))]
raw_ids = {
    str(value).strip()
    for value in raw_header
    if re.fullmatch(r"A\d{3}", str(value).strip())
}
raw_sheet_shapes = {
    ws.title: (ws.max_row - 1, ws.max_column) for ws in raw_book.worksheets
}
raw_book.close()

current = pd.read_excel(FILES["current_samples"])
current_ids = set(current.iloc[:, 0].dropna().astype(str).str.strip())
neural_ids = set(df["sample_id"].unique())

source_sets = {
    "neural": neural_ids,
    "matrix_fold_change": matrix_ids,
    "species_mapping": species_ids,
    "16S_tree": tree_ids,
    "raw_LCMS_all_sheet": raw_ids,
    "current_samples": current_ids,
}
coverage = pd.DataFrame(
    {
        "n_Axxx": {name: len(ids) for name, ids in source_sets.items()},
        "neural_missing_from_source": {
            name: len(neural_ids - ids) for name, ids in source_sets.items()
        },
    }
)
display(coverage)
display(pd.Series(raw_sheet_shapes, name="(data_rows, columns)").to_frame())
print("matrix shape:", matrix.shape, "| matrix NaN cells:", int(matrix.isna().sum().sum()))
print("matrix ↔ species set differences:", sorted(matrix_ids ^ species_ids))
print("tree ↔ species set differences:", sorted(tree_ids ^ species_ids))
print("raw-only Axxx relative to matrix:", sorted(raw_ids - matrix_ids))'''
)

md(
    """### Observation

106 个神经样本全部能在 fold-change matrix、种属表、16S tree、原始 LCMS 和 `current_samples` 中找到。`matrix.xlsx` 是 299 samples × 380 metabolites，数值区无 NaN；它与 species mapping 和 tree 的 299 个 Axxx 集合完全一致。原始 LCMS `all` sheet 有 380 个数据行和 319 个 Axxx columns，多出的 20 个样本为 A050–A058、A088–A092、A225、A250、A306、A316–A318。它们为何未进入 299-sample matrix 需要查看 LCMS/QC 生成记录。"""
)

md("""### Question

106 个细菌目前各有多少 trial、多少独立 recording、跨多少日期？""")

code(
    r'''sample_recordings = trials[[*recording_key, "sample_id"]].drop_duplicates()
per_sample = trials.groupby("sample_id").agg(
    n_trials=("segment_index", "size"),
    n_dates=("date", "nunique"),
    dates=("date", lambda x: ", ".join(sorted(x.unique()))),
)
per_sample["n_recordings"] = sample_recordings.groupby("sample_id").size()
per_sample = per_sample[["n_trials", "n_recordings", "n_dates", "dates"]]

display(per_sample[["n_trials", "n_recordings", "n_dates"]].describe().round(2))
display(pd.concat({"lowest trial count": per_sample.nsmallest(10, "n_trials"),
                   "highest trial count": per_sample.nlargest(10, "n_trials")}))

fig, axes = plt.subplots(1, 2, figsize=(11, 3.8))
sns.histplot(per_sample["n_trials"], discrete=True, ax=axes[0])
axes[0].set(title="Trials per bacterium", xlabel="operational trials")
sns.scatterplot(data=per_sample, x="n_recordings", y="n_trials", hue="n_dates", palette="viridis", ax=axes[1])
axes[1].set(title="Trial count and recording coverage", xlabel="recordings", ylabel="trials")
plt.tight_layout()
plt.show()'''
)

md(
    """### Observation

每个细菌有 17–59 个 trials（中位数 22），覆盖 4–14 个 recordings。106 个中有 100 个只在一个日期出现；跨日期样本因为增加了 recording，通常也有更多 trials。后续任何 bacteria-level 比较都不能把 2,696 个 trials 当成完全独立重复。"""
)

md("""### Question

每天测了哪些细菌？哪些细菌跨日期重复？""")

code(
    r'''date_summary = trials.groupby("date").agg(
    n_trials=("segment_index", "size"),
    n_samples=("sample_id", "nunique"),
    n_worm_recordings=("worm_key", "nunique"),
)
samples_per_date = trials.groupby("date")["sample_id"].agg(
    lambda x: ", ".join(sorted(x.unique()))
)
date_table = date_summary.join(samples_per_date.rename("sample_ids"))

cross_date = per_sample.loc[per_sample["n_dates"] > 1].sort_index()
display(date_table)
display(cross_date)

fig, ax = plt.subplots(figsize=(10, 3.6))
date_summary[["n_trials", "n_samples", "n_worm_recordings"]].plot.bar(ax=ax)
ax.set(title="Measurement coverage by date", xlabel="date", ylabel="count")
ax.tick_params(axis="x", rotation=45)
plt.tight_layout()
plt.show()'''
)

md(
    """### Observation

数据覆盖 9 个日期。每天有 11–13 个细菌、4–7 个 `(date, worm)` recordings 和 239–413 个 trials。跨日期重复仅有 6 个样本：A011、A013、A014、A024、A025、A044；每个都出现在两个日期。日期与细菌集合高度绑定，因此 date/batch 与 bacteria identity 存在明显混杂，后续必须保留日期层级。"""
)

md("""### Question

每个 trial 有哪些 neuron？neuron panel 在同一 recording 内是否稳定？""")

code(
    r'''trial_neuron_sets = traces.groupby(trial_key)["neuron"].agg(lambda x: tuple(sorted(x)))
trial_neuron_n = trial_neuron_sets.str.len()
recording_neuron_sets = traces.groupby(recording_key)["neuron"].agg(lambda x: tuple(sorted(x.unique())))
sets_per_recording = trial_neuron_sets.groupby(level=recording_key).nunique()

display(trial_neuron_n.describe().round(2).to_frame("neurons_per_trial"))
display(trial_neuron_n.value_counts().sort_index().rename("n_trials").to_frame())
print("unique neuron-set count within each recording:")
display(sets_per_recording.value_counts().sort_index().rename("n_recordings").to_frame())

recording_presence = (
    traces[[*recording_key, "neuron"]]
    .drop_duplicates()
    .assign(present=1)
    .pivot_table(index=recording_key, columns="neuron", values="present", fill_value=0)
)
fig, ax = plt.subplots(figsize=(13, 10))
sns.heatmap(recording_presence, cmap=["#f1f1f1", "#2c7fb8"], cbar=False, ax=ax)
ax.set_title("Neuron presence by recording (date × worm)")
ax.set_xlabel("neuron")
ax.set_ylabel("recording")
plt.tight_layout()
plt.show()'''
)

md(
    """### Observation

每个 trial 有 8–25 个 neurons（中位数 19）。49 个 recordings 中，每一个 recording 内的 neuron set 都完全不变；trial 之间 neuron 数差异来自 recording-level panel 不同，而不是同一 recording 中某些 trial 随机掉 neuron。全局 28-neuron union 不应被当作每个 trial 的预期 panel。"""
)

md("""### Question

数值缺失、timepoint 缺失和 neuron 缺失分别是多少？""")

code(
    r'''recording_union = traces.groupby(recording_key)["neuron"].agg(set)
recording_trial_n = trials.groupby(recording_key).size()
expected_trace_slots = sum(
    len(recording_union.loc[key]) * int(recording_trial_n.loc[key])
    for key in recording_union.index
)

missing_summary = pd.Series(
    {
        "NaN rate among stored ΔF/F0 values": df["delta_F_over_F0"].isna().mean(),
        "non-finite rate among stored ΔF/F0 values": (~np.isfinite(df["delta_F_over_F0"].to_numpy())).mean(),
        "incomplete trace rate (expected timepoints 0–44)": (trace_time["nunique"] != 45).mean(),
        "duplicate row rate under full key": df.duplicated(row_key).mean(),
        "missing neuron-trace rate within recording-specific panels": 1 - len(traces) / expected_trace_slots,
    },
    name="rate",
)
display(missing_summary.map(lambda x: f"{x:.3%}").to_frame())

neuron_coverage = traces.groupby("neuron").agg(
    n_traces=("segment_index", "size"),
    n_recordings=("date", "size"),
)
neuron_coverage["n_recordings"] = (
    traces[[*recording_key, "neuron"]].drop_duplicates().groupby("neuron").size()
)
neuron_coverage["fraction_of_all_trials"] = neuron_coverage["n_traces"] / len(trials)
display(neuron_coverage.sort_values("fraction_of_all_trials").round(3))'''
)

md(
    """### Observation

在正确 key 和 recording-specific neuron panel 下，五项 missing/duplicate 指标都是 0%。但这不表示每个 neuron 覆盖所有 recordings：例如少见的 BAGL/BAGR/CEPR/URXR 只出现在少数 recording panel。应把这种差异称为 **coverage**，而不是把未成像/未标注的 neuron 自动填成 0 或 NaN。"""
)

md("""### Question

`start_time` / `end_time` 对应哪些 timepoints？当前文件能否确认秒数和端点包含规则？""")

code(
    r'''time_contract = pd.Series(
    {
        "time_point_min": df["time_point"].min(),
        "time_point_max": df["time_point"].max(),
        "n_time_points": df["time_point"].nunique(),
        "unique_start_time": df["start_time"].nunique(),
        "start_time_value": df["start_time"].iloc[0],
        "unique_end_time": df["end_time"].nunique(),
        "end_time_value": df["end_time"].iloc[0],
        "sampling_interval_column_present": any(
            name in df.columns for name in ["sampling_interval", "dt", "time_seconds"]
        ),
    },
    name="value",
)
display(time_contract.to_frame())

ONSET = int(df["start_time"].iloc[0])
OFFSET = int(df["end_time"].iloc[0])

# 工作假设：end_time 是 exclusive boundary。修改这一行即可测试另一种约定。
END_IS_EXCLUSIVE = True
phase_table = pd.DataFrame(
    {
        "phase": ["baseline", "stimulus (working assumption)", "post-stimulus"],
        "timepoints": [
            list(range(0, ONSET)),
            list(range(ONSET, OFFSET if END_IS_EXCLUSIVE else OFFSET + 1)),
            list(range(OFFSET if END_IS_EXCLUSIVE else OFFSET + 1, 45)),
        ],
    }
)
phase_table["n_points"] = phase_table["timepoints"].str.len()
display(phase_table)'''
)

md(
    """### Observation

所有行的 onset marker 都是 5，offset marker 都是 15。表本身没有采样间隔或绝对秒数列，因此只能确认索引，不能独立证明 1 point = 1 s。当前工作约定把 `end_time` 视为 exclusive boundary：baseline 0–4（5 点）、stimulus 5–14（10 点）、post 15–44（30 点）。若 protocol 把 15 也算刺激期，只需切换 `END_IS_EXCLUSIVE`；在 protocol 确认前不要把这项约定写成事实。"""
)

md("""### Question

能否从已归一化的 traces 反推当前 baseline 使用了哪些点？""")

code(
    r'''baseline_rows = df[df["time_point"] < ONSET]
baseline_trace = baseline_rows.groupby(trace_key)["delta_F_over_F0"].agg(
    baseline_mean="mean",
    baseline_sd="std",
    baseline_range=lambda x: x.max() - x.min(),
)

display(
    baseline_trace[["baseline_mean", "baseline_sd", "baseline_range"]]
    .describe(percentiles=[0.5, 0.9, 0.95, 0.99])
    .round(5)
)
print("max |mean ΔF/F0 over timepoints 0–4|:", baseline_trace["baseline_mean"].abs().max())'''
)

md(
    """### Observation

48,701 条 trace 在 timepoints 0–4 上的平均 $\Delta F/F_0$ 都等于 0（最大浮点误差约 $3.6\times10^{-16}$）。这强烈支持 0–4 是归一化所用 baseline window，并且归一化使该窗口均值归零。仅凭归一化结果仍不能证明上游 $F_0$ 的精确公式、去漂白步骤或是否有其他预处理；这些需要 raw fluorescence 或生成记录。"""
)

md("""### Question

每个 trial 的 baseline fluctuation 多大，是否集中在某些日期或 recordings？""")

code(
    r'''trial_baseline = baseline_trace.groupby(level=trial_key).agg(
    n_neurons=("baseline_sd", "size"),
    median_neuron_sd=("baseline_sd", "median"),
    p90_neuron_sd=("baseline_sd", lambda x: x.quantile(0.9)),
    max_neuron_sd=("baseline_sd", "max"),
)

display(
    trial_baseline.describe(percentiles=[0.5, 0.9, 0.95, 0.99]).round(4)
)
display(trial_baseline.nlargest(12, "median_neuron_sd").round(4))

plot_baseline = trial_baseline.reset_index()
fig, axes = plt.subplots(1, 2, figsize=(13, 4))
sns.histplot(plot_baseline["median_neuron_sd"], bins=40, ax=axes[0])
axes[0].axvline(plot_baseline["median_neuron_sd"].quantile(0.99), color="C3", ls="--", label="99th percentile")
axes[0].legend()
axes[0].set(title="Trial baseline fluctuation", xlabel="median neuron baseline SD")
sns.boxplot(data=plot_baseline, x="date", y="median_neuron_sd", color="#8ecae6", ax=axes[1])
axes[1].tick_params(axis="x", rotation=45)
axes[1].set(title="Baseline fluctuation by date", xlabel="date", ylabel="median neuron baseline SD")
plt.tight_layout()
plt.show()'''
)

md(
    """### Observation

trial-level median neuron baseline SD 的中位数为 0.0523，95th percentile 为 0.0960，99th percentile 为 0.1212，最大值 0.2571。最高的 6 个 trial 全来自 2026-03-31 的 w4，说明至少存在 recording-level 状态或质量差异。每条 baseline 只有 5 个点，SD 估计本身很不稳定；这里把它当作排序和看原始曲线的诊断量，不设自动剔除阈值。"""
)

md("""### Question

baseline SD 最大的 traces 是单点跳变，还是持续波动？""")

code(
    r'''worst_baseline = baseline_trace.nlargest(8, "baseline_sd").reset_index()
display(worst_baseline.round(4))

worst_baseline_curves = df.merge(worst_baseline[trace_key], on=trace_key, how="inner")
g = sns.FacetGrid(
    worst_baseline_curves,
    col="neuron",
    col_wrap=4,
    hue="sample_id",
    sharey=False,
    height=2.5,
    aspect=1.25,
)
g.map_dataframe(sns.lineplot, x="time_point", y="delta_F_over_F0", estimator=None)
for ax in g.axes.flat:
    ax.axvspan(0, ONSET - 1, color="#bbbbbb", alpha=0.18)
    ax.axvline(ONSET, color="C2", ls="--", lw=1)
    ax.axvline(OFFSET, color="C3", ls="--", lw=1)
g.set_axis_labels("time_point", "ΔF/F0")
g.fig.suptitle("Traces with the largest baseline SD", y=1.03)
plt.show()'''
)

md(
    """### Observation

最大的单条 baseline SD 为 1.693（2026-04-14, w2, A013, segment 19, ASJR）。高 baseline SD 涉及 ASJR、ASJL、ASIL、ASEL、AWCON 等不同 neurons 和多个 recordings；原始曲线可用于区分宽幅波动、漂移和个别点跳变。现阶段只标记、追溯，不删除。"""
)

md("""### Question

是否存在极端 $\Delta F/F_0$ traces，它们集中在哪些 neurons？""")

code(
    r'''trace_stats = df.groupby(trace_key)["delta_F_over_F0"].agg(
    trace_min="min",
    trace_max="max",
    trace_mean="mean",
    trace_sd="std",
)
trace_stats["max_abs"] = trace_stats[["trace_min", "trace_max"]].abs().max(axis=1)
trace_stats = trace_stats.join(baseline_trace)

quantiles = trace_stats["max_abs"].quantile([0.5, 0.9, 0.95, 0.99, 0.999, 1.0])
threshold_counts = []
for threshold in [2, 3, 5, 8]:
    selected = trace_stats[trace_stats["max_abs"] >= threshold].reset_index()
    threshold_counts.append(
        {
            "|ΔF/F0| threshold": threshold,
            "n_traces": len(selected),
            "AWCON_traces": int((selected["neuron"] == "AWCON").sum()),
        }
    )

top_extreme = trace_stats.nlargest(15, "max_abs").reset_index()
display(quantiles.rename("max_abs").to_frame().round(4))
display(pd.DataFrame(threshold_counts))
display(top_extreme.round(4))

fig, ax = plt.subplots(figsize=(8, 3.8))
sns.histplot(trace_stats["max_abs"], bins=100, ax=ax)
ax.set_yscale("log")
ax.axvline(quantiles.loc[0.99], color="C3", ls="--", label="99th percentile")
ax.set(title="Distribution of per-trace maximum |ΔF/F0|", xlabel="maximum |ΔF/F0|", ylabel="trace count (log scale)")
ax.legend()
plt.tight_layout()
plt.show()'''
)

md(
    """### Observation

per-trace 最大绝对幅度的 99th percentile 为 2.479，99.9th percentile 为 7.548，最大值为 10.028。共有 118 条 trace 达到 $|\Delta F/F_0|\ge5$，其中 94 条是 AWCON；达到 8 的 35 条中 33 条是 AWCON。极值明显有 neuron 结构，不能用单一全局阈值自动当作坏数据。"""
)

md("""### Question

幅度最大的 traces 看起来像孤立坏点，还是完整的刺激相关动态？""")

code(
    r'''top_plot_keys = top_extreme.head(12)[trace_key]
top_curves = df.merge(top_plot_keys, on=trace_key, how="inner")
top_curves["trace_label"] = (
    top_curves["date"].astype(str) + " "
    + top_curves["worm_key"] + " "
    + top_curves["sample_id"] + " seg="
    + top_curves["segment_index"].astype(str) + " "
    + top_curves["neuron"]
)

g = sns.FacetGrid(
    top_curves,
    col="trace_label",
    col_wrap=3,
    sharey=False,
    height=2.6,
    aspect=1.25,
)
g.map_dataframe(sns.lineplot, x="time_point", y="delta_F_over_F0", estimator=None, color="#264653")
for ax in g.axes.flat:
    ax.axvspan(ONSET, OFFSET, color="#f4a261", alpha=0.18)
    ax.axvline(ONSET, color="C2", ls="--", lw=1)
    ax.axvline(OFFSET, color="C3", ls="--", lw=1)
g.set_titles("{col_name}", size=8)
g.set_axis_labels("time_point", "ΔF/F0")
g.fig.suptitle("Largest-amplitude traces", y=1.02)
plt.show()'''
)

md(
    """### Observation

最高幅 trace（2026-04-10, w6, A289, segment 40, AWCON）从 onset 后连续上升，在 timepoint 11 达到 10.028，随后缓慢衰减到记录末端；它不是一个孤立数值坏点。多个高幅 AWCON traces 也呈连续动力学。它们可能是真实强响应，也可能受很小 $F_0$、饱和或上游处理影响；需要 raw fluorescence、曝光/增益和 $F_0$ 记录才能判定。当前只建立 review list，不剔除。"""
)

md(
    """## 本轮 measurement checkpoint

已经可以直接从文件确认：

- 106 个神经刺激样本，49 个 date × worm recordings，2,696 个 operational trials，48,701 条完整 45-point traces。
- `(date, worm_key)` 是最低限度的 recording key；忽略 date 会制造伪重复。
- neuron panel 在每个 recording 内恒定，跨 recording 不同。
- 存储值、timepoints、正确复合键和 recording-specific neuron slots 均无 missing/duplicate。
- onset/offset markers 为 5/15；0–4 的 baseline 均值对每条 trace 精确归零。
- baseline fluctuation 存在 recording-level 集中；极端幅度主要集中在 AWCON，并呈完整连续动态。

## 进入下一轮前需要确认的 protocol 问题

1. 每个 `time_point` 是否严格等于 1 s？一个 point 是一帧还是一个完整 volume？
2. `end_time=15` 是 exclusive boundary 还是刺激仍覆盖 timepoint 15？
3. `(date, worm_key)` 是否一一对应独立动物？是否存在同一动物跨日/跨 recording？
4. `segment_index` 的来源是什么？它是否严格等于一次 stimulus presentation/trial？
5. $\Delta F/F_0$ 的 $F_0$ 精确公式、motion correction、photobleaching correction 和去噪步骤是什么？
6. 是否保存 raw fluorescence、stimulus TTL/同步日志、volume timing、voxel size 和 z-step？
7. 原始 LCMS 中 20 个 raw-only samples 为什么未进入 299-sample fold-change matrix？

这些问题确认前，不继续做跨菌株表征、相似性或 feature screen。"""
)

nb.cells = cells
nb.metadata.kernelspec = {
    "display_name": "Python 3 (pixi: bacteria_analysis)",
    "language": "python",
    "name": "python3",
}
nb.metadata.language_info = {"name": "python", "version": "3.11"}
nbf.write(nb, OUT)
print(f"Wrote {OUT} with {len(cells)} cells")

