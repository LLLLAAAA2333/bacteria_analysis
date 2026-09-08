from pathlib import Path
import json
import re

import numpy as np
import pandas as pd
from Bio import Phylo
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data"


def scalar(value):
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return None if pd.isna(value) else float(value)
    if isinstance(value, pd.Timestamp):
        return value.isoformat()
    return value


df = pd.read_parquet(DATA / "106bac.parquet")
df["sample_id"] = df["stim_name"].str.extract(r"^(A\d{3})", expand=False)

trace_key = ["worm_key", "stimulus", "segment_index", "neuron"]
trial_key = ["worm_key", "stimulus", "segment_index"]
row_key = trace_key + ["time_point"]

trace_time = df.groupby(trace_key, observed=True)["time_point"].agg(["size", "nunique", "min", "max"])
trials = df[trial_key + ["sample_id", "date", "stim_name"]].drop_duplicates()
trial_neurons = df.groupby(trial_key, observed=True)["neuron"].agg(lambda x: tuple(sorted(x.unique())))

summary = {
    "shape": list(df.shape),
    "columns": list(df.columns),
    "null_counts": {k: int(v) for k, v in df.isna().sum().items()},
    "unique": {c: int(df[c].nunique(dropna=False)) for c in df.columns},
    "dates": sorted(df["date"].astype(str).unique().tolist()),
    "worms": sorted(df["worm_key"].unique().tolist()),
    "neurons": sorted(df["neuron"].unique().tolist()),
    "time_points": sorted(df["time_point"].unique().tolist()),
    "start_end_pairs": df[["start_time", "end_time"]].drop_duplicates().values.tolist(),
    "duplicate_row_keys": int(df.duplicated(row_key).sum()),
    "n_trials": int(len(trials)),
    "n_traces": int(len(trace_time)),
    "trace_time_shapes": (
        trace_time.value_counts().rename("n_traces").reset_index().to_dict(orient="records")
    ),
    "stimulus_to_sample_id_max_nunique": int(df.groupby("stimulus")["sample_id"].nunique().max()),
    "sample_id_to_stimulus_max_nunique": int(df.groupby("sample_id")["stimulus"].nunique().max()),
    "stim_name_examples": sorted(df["stim_name"].unique().tolist())[:8],
    "stim_name_suffixes": df["stim_name"].str.replace(r"^A\d{3}\s*", "", regex=True).value_counts().to_dict(),
    "stimuli_per_worm_segment_distribution": (
        df.groupby(["worm_key", "segment_index"])["stimulus"].nunique().value_counts().sort_index().to_dict()
    ),
    "dates_per_worm_segment_distribution": (
        df.groupby(["worm_key", "segment_index"])["date"].nunique().value_counts().sort_index().to_dict()
    ),
    "neurons_per_trial_summary": {
        k: scalar(v) for k, v in df.groupby(trial_key)["neuron"].nunique().describe().items()
    },
    "neurons_per_trial_distribution": (
        df.groupby(trial_key)["neuron"].nunique().value_counts().sort_index().to_dict()
    ),
}

per_sample = trials.groupby("sample_id").agg(
    n_trials=("segment_index", "size"),
    n_worms=("worm_key", "nunique"),
    n_dates=("date", "nunique"),
    dates=("date", lambda x: tuple(sorted(x.unique()))),
)
summary["per_sample_describe"] = {
    c: {k: scalar(v) for k, v in per_sample[c].describe().items()}
    for c in ["n_trials", "n_worms", "n_dates"]
}
summary["n_cross_date_samples"] = int((per_sample["n_dates"] > 1).sum())
summary["cross_date_samples"] = {
    k: list(v) for k, v in per_sample.loc[per_sample["n_dates"] > 1, "dates"].items()
}
summary["trials_per_date"] = {
    str(k): int(v) for k, v in trials.groupby("date").size().items()
}
summary["samples_per_date"] = {
    str(k): sorted(v.unique().tolist()) for k, v in trials.groupby("date")["sample_id"]
}
summary["trials_per_worm"] = {
    str(k): int(v) for k, v in trials.groupby("worm_key").size().items()
}

# Existing rows, expected time grid, and structural neuron coverage answer different missingness questions.
expected_time = set(df["time_point"].unique())
summary["value_missing_rate"] = float(df["delta_F_over_F0"].isna().mean())
summary["nonfinite_count"] = int((~np.isfinite(df["delta_F_over_F0"].to_numpy())).sum())
summary["incomplete_trace_count"] = int((trace_time["nunique"] != len(expected_time)).sum())

all_trials = pd.MultiIndex.from_frame(trials[trial_key])
present = df[trial_key + ["neuron"]].drop_duplicates()
neuron_present_n = present.groupby("neuron").size()
neuron_missing = (1 - neuron_present_n / len(all_trials)).sort_values(ascending=False)
summary["neuron_structural_missing_rate"] = {
    k: float(v) for k, v in neuron_missing.items()
}

onset = int(df["start_time"].iloc[0])
baseline = df[df["time_point"] < onset]
baseline_trace = baseline.groupby(trace_key)["delta_F_over_F0"].agg(
    baseline_mean="mean",
    baseline_sd="std",
    baseline_range=lambda x: x.max() - x.min(),
)
trial_baseline = baseline_trace.groupby(level=trial_key).agg(
    n_neurons=("baseline_sd", "size"),
    median_neuron_sd=("baseline_sd", "median"),
    p90_neuron_sd=("baseline_sd", lambda x: x.quantile(0.9)),
    max_neuron_sd=("baseline_sd", "max"),
    median_abs_baseline_mean=("baseline_mean", lambda x: x.abs().median()),
)
for c in trial_baseline.columns:
    summary[f"trial_baseline_{c}_quantiles"] = {
        str(q): scalar(v)
        for q, v in trial_baseline[c].quantile([0, 0.25, 0.5, 0.75, 0.9, 0.95, 0.99, 1]).items()
    }

trace_stats = df.groupby(trace_key)["delta_F_over_F0"].agg(
    trace_min="min", trace_max="max", trace_mean="mean", trace_sd="std"
)
trace_stats["max_abs"] = trace_stats[["trace_min", "trace_max"]].abs().max(axis=1)
trace_stats = trace_stats.join(baseline_trace)
summary["trace_max_abs_quantiles"] = {
    str(q): scalar(v)
    for q, v in trace_stats["max_abs"].quantile([0.5, 0.9, 0.95, 0.99, 0.999, 1]).items()
}
summary["top_extreme_traces"] = (
    trace_stats.nlargest(12, "max_abs").reset_index().to_dict(orient="records")
)
summary["top_baseline_fluctuation_traces"] = (
    trace_stats.nlargest(12, "baseline_sd").reset_index().to_dict(orient="records")
)

# Cross-modality sample coverage.
matrix = pd.read_excel(DATA / "matrix.xlsx", index_col=0)
matrix.index = matrix.index.astype(str).str.strip()
matrix_ids = {x for x in matrix.index if re.fullmatch(r"A\d{3}", x)}

species = pd.read_excel(DATA / "GM300_bacteria_species_summary.xlsx", sheet_name="Axxx_species_mapping")
species_ids = set(species["AID"].dropna().astype(str).str.strip())

tree = Phylo.read(DATA / "16S.aln.trim.fa.treefile", "newick")
tree_ids = {t.name for t in tree.get_terminals() if t.name}
neural_ids = set(df["sample_id"].dropna().unique())

current = pd.read_excel(DATA / "current_samples.xlsx")
current_ids = set(current.iloc[:, 0].dropna().astype(str).str.strip())

summary["source_coverage"] = {
    "neural": len(neural_ids),
    "matrix": len(matrix_ids),
    "species": len(species_ids),
    "tree": len(tree_ids),
    "current_samples": len(current_ids),
    "neural_missing_from_matrix": sorted(neural_ids - matrix_ids),
    "neural_missing_from_species": sorted(neural_ids - species_ids),
    "neural_missing_from_tree": sorted(neural_ids - tree_ids),
    "neural_missing_from_current_samples": sorted(neural_ids - current_ids),
}
summary["matrix_shape"] = list(matrix.shape)
summary["matrix_duplicate_ids"] = int(matrix.index.duplicated().sum())
summary["matrix_value_missing_rate"] = float(matrix.isna().to_numpy().mean())

for path in [DATA / "metabolism_raw_data.xlsx"]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets = {}
    for worksheet in workbook.worksheets:
        header = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
        sample_columns = [str(v).strip() for v in header if re.fullmatch(r"A\d{3}", str(v).strip())]
        sheets[worksheet.title] = {
            "max_row": worksheet.max_row,
            "max_column": worksheet.max_column,
            "first_20_headers": header[:20],
            "n_Axxx_columns": len(sample_columns),
            "first_Axxx_columns": sample_columns[:5],
            "last_Axxx_columns": sample_columns[-5:],
        }
    summary["raw_lcms_sheets"] = sheets
    workbook.close()

print(json.dumps(summary, ensure_ascii=False, indent=2, default=scalar))

