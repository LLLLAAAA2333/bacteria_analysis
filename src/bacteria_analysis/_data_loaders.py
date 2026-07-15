"""Low-level data loading helpers for metabolite matrix and stimulus mapping."""

from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path
import re

import pandas as pd
from openpyxl import load_workbook

METABOLITE_NAME_CANONICAL_OVERRIDES = {
    "Beta-Muricholic acid (β-MCA)": "Beta-Muricholic acid (beta-MCA)",
    "Tauro-α-muricholic acid (α-TMCA)": "Tauro-alpha-muricholic acid (alpha-TMCA)",
    "Tauro-β-muricholic acid (β-TMCA)": "Tauro-beta-muricholic acid (beta-TMCA)",
    "Tauro-α-muricholic acid (ω-TMCA)": "Tauro-omega-muricholic acid (omega-TMCA)",
    "Tauro-alpha-muricholic acid (omega-TMCA)": "Tauro-omega-muricholic acid (omega-TMCA)",
    "Omega-muricholic acid (ω-MCA)": "Omega-muricholic acid (omega-MCA)",
    "α-Ketoglutaric acid": "Alpha-Ketoglutaric acid",
    "Riboflavin-5′-monophosphate": "Riboflavin-5'-monophosphate",
    "Adenosine-5′-triphosphate(ATP)": "Adenosine-5'-triphosphate(ATP)",
    "Adenosine-5′-diphosphate(ADP)": "Adenosine-5'-diphosphate(ADP)",
    "Adenosine-3′,5′-cyclic monophosphate(cAMP)": "Adenosine-3',5'-cyclic monophosphate(cAMP)",
}

METABOLITE_TEXT_REPLACEMENTS: tuple[tuple[str, str], ...] = (
    ("（", "("),
    ("）", ")"),
    ("α", "alpha"),
    ("β", "beta"),
    ("γ", "gamma"),
    ("δ", "delta"),
    ("ω", "omega"),
    ("¦Á", "alpha"),
    ("¦Â", "beta"),
    ("¦Ã", "gamma"),
    ("¦Ä", "delta"),
    ("¦Ø", "omega"),
    ("′", "'"),
    ("’", "'"),
    ("‘", "'"),
    ("ʼ", "'"),
    ("ʹ", "'"),
    ("ˈ", "'"),
    ("´", "'"),
    ("`", "'"),
    ("¡ä", "'"),
)


def read_metabolite_matrix(path: str | Path) -> pd.DataFrame:
    """Load a metabolite matrix workbook into a sample_id-indexed DataFrame."""

    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        first_sheet = workbook.worksheets[0]
        header_row = next(first_sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    finally:
        workbook.close()

    header_values = [_canonicalize_metabolite_name(value) for value in header_row]
    if header_values and pd.Index(header_values).duplicated().any():
        raise ValueError("metabolite column names must be unique")

    frame = pd.read_excel(path, engine="openpyxl", sheet_name=0)
    return _normalize_matrix_frame(frame)


def build_stimulus_sample_map(metadata: pd.DataFrame, *, matrix_sample_ids: pd.Index) -> pd.DataFrame:
    """Build stimulus-to-sample mapping from trial metadata and matrix columns."""

    _require_columns(metadata, ("stimulus", "stim_name"), "trial metadata")
    rows = metadata.loc[:, ["stimulus", "stim_name"]].copy()
    rows["stimulus"] = rows["stimulus"].fillna("").astype(str).str.strip()
    rows["stim_name"] = rows["stim_name"].fillna("").astype(str).str.strip()
    if (rows["stimulus"] == "").any():
        raise ValueError("stimulus values must be non-empty")
    if (rows["stim_name"] == "").any():
        raise ValueError("stim_name values must be non-empty")

    stimulus_name_counts = rows.groupby("stimulus", sort=False)["stim_name"].nunique(dropna=False)
    conflicting = stimulus_name_counts.loc[stimulus_name_counts > 1]
    if not conflicting.empty:
        raise ValueError(
            "stimulus values must map to exactly one stim_name: "
            + ", ".join(conflicting.index.astype(str).tolist())
        )

    collapsed = rows.drop_duplicates(subset=["stimulus"], keep="first").copy()
    collapsed["sample_id"] = collapsed["stim_name"].map(_sample_id_from_stim_name)
    if (collapsed["sample_id"] == "").any():
        raise ValueError("sample_id values derived from stim_name must be non-empty")
    if collapsed["sample_id"].duplicated().any():
        duplicates = collapsed.loc[collapsed["sample_id"].duplicated(keep=False), "sample_id"].tolist()
        raise ValueError(f"derived sample_id values must be unique: {', '.join(dict.fromkeys(duplicates))}")

    matrix_sample_set = set(matrix_sample_ids.astype(str))
    missing = sorted(set(collapsed["sample_id"]).difference(matrix_sample_set))
    if missing:
        raise ValueError(f"derived sample_id values must exist in the matrix: {', '.join(missing)}")

    return collapsed.loc[:, ["stimulus", "stim_name", "sample_id"]].reset_index(drop=True)


def _normalize_matrix_frame(frame: pd.DataFrame) -> pd.DataFrame:
    normalized = frame.copy()
    if normalized.empty:
        return normalized

    if "sample_id" in normalized.columns:
        sample_column = "sample_id"
    else:
        sample_column = normalized.columns[0]

    sample_ids = normalized[sample_column]
    if sample_ids.isna().any():
        raise ValueError("sample_id values must be non-empty")

    sample_ids = sample_ids.astype(str).str.strip()
    if (sample_ids == "").any():
        raise ValueError("sample_id values must be non-empty")
    if sample_ids.duplicated().any():
        raise ValueError("sample_id values must be unique")

    metabolite_columns = [c for c in normalized.columns if c != sample_column]
    metabolite_column_names = [_canonicalize_metabolite_name(c) for c in metabolite_columns]
    if any(not name or name.lower().startswith("unnamed:") for name in metabolite_column_names):
        raise ValueError("metabolite column names must be non-empty")
    if pd.Index(metabolite_column_names).duplicated().any():
        raise ValueError("metabolite column names must be unique")

    name_iter = iter(metabolite_column_names)
    normalized.columns = [c if c == sample_column else next(name_iter) for c in normalized.columns]
    normalized[sample_column] = sample_ids
    normalized = normalized.set_index(sample_column)
    normalized.index.name = "sample_id"
    return normalized


def _sample_id_from_stim_name(stim_name: str) -> str:
    parts = str(stim_name).strip().split()
    return parts[0].strip() if parts else ""


def _canonicalize_metabolite_name(value: object) -> str:
    if value is None:
        return ""
    normalized = str(value).strip()
    direct_override = METABOLITE_NAME_CANONICAL_OVERRIDES.get(normalized)
    if direct_override is not None:
        return direct_override
    for original, replacement in METABOLITE_TEXT_REPLACEMENTS:
        if original in normalized:
            normalized = normalized.replace(original, replacement)
    normalized = re.sub(r"\s+", " ", normalized).strip()
    return METABOLITE_NAME_CANONICAL_OVERRIDES.get(normalized, normalized)


def enrich_neural_dataframe(
    df: pd.DataFrame | str | Path,
    species_path: str | Path = "data/GM300_bacteria_species_summary.xlsx",
) -> pd.DataFrame:
    """Add species and genus columns by joining stim_name → AID → GM300 metadata.

    Parses ``stim_name`` (e.g. "A001 stationary") to extract the AID,
    then looks up ``species_clean`` and ``genus_clean`` from the GM300
    species summary spreadsheet.

    Returns a copy of the DataFrame with additional columns:
    ``species``, ``genus``, and ``aid`` (the parsed AID).
    """

    if isinstance(df, (str, Path)):
        df = pd.read_parquet(df)

    species = pd.read_excel(Path(species_path), engine="openpyxl")

    out = df.copy()
    out["aid"] = out["stim_name"].fillna("").astype(str).str.strip().str.split().str[0]
    aid_map = species.set_index("AID")[["species_clean", "genus_clean"]]
    aid_map.columns = ["species", "genus"]

    return out.join(aid_map, on="aid")


def _require_columns(frame: pd.DataFrame, required: Iterable[str], label: str) -> pd.DataFrame:
    missing = [c for c in required if c not in frame.columns]
    if missing:
        raise ValueError(f"{label} must include columns: {', '.join(missing)}")
    return frame.copy()
